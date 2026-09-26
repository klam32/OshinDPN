import io
import json
import os
import tempfile
import time
import unittest
import uuid
from pathlib import Path
from urllib.parse import parse_qs, urlparse
from unittest.mock import MagicMock, patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.main import app, attempts
from backend.db import connect, services
from backend.logic import deliver_mail

HEADERS = {'X-Requested-With': 'OshinWeb', 'Origin': 'http://localhost:5173'}

class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.environment = patch.dict(os.environ, {'DATABASE_PATH': str(Path(self.directory.name) / 'test.db'), 'ADMIN_EMAIL': 'admin@example.com', 'ADMIN_PASSWORD': 'Admin-test-strong-123', 'AUTH_SECRET': 'test-auth-secret-with-at-least-32-characters', 'GOOGLE_CLIENT_ID': 'test-google-client', 'GOOGLE_CLIENT_SECRET': 'test-google-secret', 'GOOGLE_REDIRECT_URI': 'http://testserver/api/auth/google/callback', 'FRONTEND_URL': 'http://testserver', 'LLM_NAME': 'vertex', 'PROJECT_ID': '', 'LOCATION': '', 'VERTEX_MODEL_NAME': '', 'GOOGLE_APPLICATION_CREDENTIALS': '', 'GOOGLE_APPLICATION_CREDENTIALS_JSON': '', 'GOOGLE_APPLICATION_CREDENTIALS_BASE64': '', 'SMTP_HOST': '', 'SMTP_FROM': ''})
        self.environment.start()
        attempts.clear()
        self.client = TestClient(app, headers=HEADERS)
        self.client.__enter__()
        self.admin = TestClient(app, headers=HEADERS)
        self.login_with_otp(self.admin, 'admin@example.com', 'Admin-test-strong-123')
        with connect() as c:
            c.execute("DELETE FROM outbox WHERE kind='auth_otp'")
            c.execute('DELETE FROM auth_otp')

    def login_with_otp(self, client, email, password, name=None, action='login'):
        payload = {'email': email, 'password': password}
        if name is not None:
            payload['name'] = name
        with patch('backend.main.new_otp_code', return_value='100000'), patch.dict(
            os.environ, {'SMTP_HOST': 'smtp.example.com', 'SMTP_FROM': 'service@example.com'}
        ), patch('backend.logic.smtplib.SMTP'):
            started = client.post(f'/api/auth/{action}', json=payload)
        self.assertEqual(started.status_code, 200, started.text)
        challenge = started.json()
        self.assertTrue(challenge['otp_required'])
        verified = client.post('/api/auth/otp/verify', json={
            'challenge_id': challenge['challenge_id'], 'code': '100000'
        })
        self.assertEqual(verified.status_code, 200, verified.text)
        return verified

    def tearDown(self):
        self.admin.close()
        self.client.__exit__(None, None, None)
        self.environment.stop()
        self.directory.cleanup()

    def payload(self, service_id='cleaning', mode='quote'):
        service = next(s for s in services() if s['id'] == service_id)
        details = {f['key']: max(2, f.get('min', 1)) if f['type'] == 'number' else f['options'][0] if f['type'] == 'select' else 'Thông tin kiểm thử' for f in service['fields']}
        return {'service_id': service_id, 'subtype': service['subservices'][0], 'details': details if mode == 'quote' else {}, 'mode': mode, 'name': 'Khách kiểm thử', 'phone': '0901234567', 'email': 'customer@example.com' if mode == 'quote' else '', 'address': '12 Đường kiểm thử, Cần Thơ', 'preferred_date': '', 'consent': True, 'request_id': str(uuid.uuid4())}

    def create(self, **kwargs):
        result = self.client.post('/api/orders', json=self.payload(**kwargs))
        self.assertEqual(result.status_code, 200, result.text)
        return result.json()

    def test_auth_role_and_session_rotation(self):
        self.client.get('/api/auth/me')
        old = self.client.cookies.get('oshin_session')
        with patch('backend.main.new_otp_code', return_value='100000'), patch.dict(
            os.environ, {'SMTP_HOST': 'smtp.example.com', 'SMTP_FROM': 'service@example.com'}
        ), patch('backend.logic.smtplib.SMTP'):
            started = self.client.post('/api/auth/register', json={'email': 'new@example.com', 'password': 'Password123', 'name': 'Khách mới'})
        self.assertEqual(started.status_code, 200, started.text)
        challenge = started.json()
        self.assertEqual(self.client.post('/api/auth/otp/verify', json={
            'challenge_id': challenge['challenge_id'], 'code': '999999'
        }).status_code, 401)
        result = self.client.post('/api/auth/otp/verify', json={
            'challenge_id': challenge['challenge_id'], 'code': '100000'
        })
        self.assertEqual(result.status_code, 200, result.text)
        self.assertEqual(result.json()['user']['role'], 'customer')
        self.assertNotEqual(old, self.client.cookies.get('oshin_session'))
        with connect() as c:
            otp_mail = c.execute("SELECT status,body FROM outbox WHERE kind='auth_otp' ORDER BY created DESC LIMIT 1").fetchone()
        self.assertEqual(otp_mail['status'], 'sent')
        self.assertEqual(otp_mail['body'], '')
        self.assertEqual(self.client.get('/api/admin/dashboard').status_code, 403)
        self.assertEqual(self.client.post('/api/auth/register', json={'email': 'new@example.com', 'password': 'Password123', 'name': 'Khách mới', 'role': 'admin'}).status_code, 422)
        self.client.post('/api/auth/logout')
        self.assertIsNone(self.client.get('/api/auth/me').json()['user'])

    def test_google_oauth_state_pkce_nonce_and_account_link(self):
        start = self.client.get('/api/auth/google/start', follow_redirects=False)
        self.assertEqual(start.status_code, 302)
        query = parse_qs(urlparse(start.headers['location']).query)
        self.assertEqual(query['client_id'], ['test-google-client'])
        self.assertEqual(query['scope'], ['openid email profile'])
        self.assertEqual(query['code_challenge_method'], ['S256'])
        state = query['state'][0]
        self.assertEqual(start.cookies.get('oshin_oauth_state'), state)
        for cookie in list(self.client.cookies.jar):
            if cookie.name == 'oshin_session':
                self.client.cookies.jar.clear(cookie.domain, cookie.path, cookie.name)
        with connect() as c:
            saved = dict(c.execute('SELECT * FROM oauth_states WHERE state=?', (state,)).fetchone())
        token_response = MagicMock()
        token_response.json.return_value = {'id_token': 'verified-id-token'}
        claims = {'sub': 'google-subject-1', 'email': 'google@example.com', 'email_verified': True, 'name': 'Google User', 'nonce': saved['nonce']}
        with patch('backend.main.httpx.post', return_value=token_response) as exchange, patch(
            'backend.main.google_id_token.verify_oauth2_token', return_value=claims
        ):
            callback = self.client.get(f'/api/auth/google/callback?code=test-code&state={state}', follow_redirects=False)
        self.assertEqual(callback.status_code, 303, callback.text)
        self.assertEqual(callback.headers['location'], 'http://testserver/#account')
        self.assertEqual(self.client.get('/api/auth/me').json()['user']['email'], 'google@example.com')
        self.assertEqual(exchange.call_args.kwargs['data']['code_verifier'], saved['code_verifier'])
        with connect() as c:
            self.assertIsNotNone(c.execute('SELECT 1 FROM google_identities WHERE sub=?', ('google-subject-1',)).fetchone())
            self.assertIsNone(c.execute('SELECT 1 FROM oauth_states WHERE state=?', (state,)).fetchone())
        replay = self.client.get(f'/api/auth/google/callback?code=test-code&state={state}', follow_redirects=False)
        self.assertEqual(replay.status_code, 400)

    def test_google_oauth_start_uses_canonical_host(self):
        with TestClient(app, base_url='https://preview.example', headers=HEADERS) as preview:
            start = preview.get('/api/auth/google/start', follow_redirects=False)
        self.assertEqual(start.status_code, 302)
        self.assertEqual(start.headers['location'], 'http://testserver/api/auth/google/start')

    def test_all_seven_service_forms_and_excel(self):
        for service in services():
            with self.subTest(service=service['id']):
                order = self.create(service_id=service['id'])
                self.assertGreater(order['quote']['total'], 0)
                result = self.client.get(f'/api/orders/{order["id"]}/xlsx')
                self.assertEqual(result.status_code, 200)
                sheet = load_workbook(io.BytesIO(result.content)).active
                self.assertEqual(sheet['B4'].value, 'Khách kiểm thử')
                self.assertEqual(sheet['B5'].value, '0901234567')
                self.assertTrue((Path(self.directory.name) / 'exports' / f'{order["id"]}.xlsx').exists())

    def test_survey_minimal_and_foreign_order_hidden(self):
        order = self.create(mode='survey')
        self.assertTrue(order['quote']['survey'])
        self.assertEqual(order['quote']['total'], 0)
        with TestClient(app, headers=HEADERS) as stranger:
            self.assertEqual(stranger.get(f'/api/orders/{order["id"]}/xlsx').status_code, 404)
            self.assertEqual(stranger.get('/api/orders').json(), [])

    def test_invalid_quantity_and_email(self):
        data = self.payload()
        data['details']['area'] = -1
        self.assertEqual(self.client.post('/api/orders', json=data).status_code, 422)
        data['details']['area'] = 'nan'
        self.assertEqual(self.client.post('/api/orders', json=data).status_code, 422)
        data['details']['area'] = 50
        data['email'] = ''
        self.assertEqual(self.client.post('/api/orders', json=data).status_code, 422)
        data['email'] = 'invalid'
        self.assertEqual(self.client.post('/api/orders', json=data).status_code, 422)

    def test_guest_orders_transfer_on_registration_and_idempotency(self):
        payload = self.payload()
        one = self.client.post('/api/orders', json=payload).json()
        two = self.client.post('/api/orders', json=payload).json()
        self.assertEqual(one['id'], two['id'])
        self.login_with_otp(self.client, 'owner@example.com', 'Password123', 'Chủ đơn', 'register')
        self.assertEqual(len(self.client.get('/api/orders').json()), 1)
        with TestClient(app, headers=HEADERS) as another:
            self.login_with_otp(another, 'owner@example.com', 'Password123')
            self.assertEqual(another.get('/api/orders').json()[0]['id'], one['id'])

    def test_chat_handoff_presence_expiry_and_ai_fallback(self):
        self.client.get('/api/chat')
        r = self.client.post('/api/chat/send', json={'content': 'Tôi cần dịch vụ vệ sinh', 'mode': 'ai'}).json()
        self.assertEqual(r['messages'][-1]['role'], 'assistant')
        self.admin.post('/api/admin/presence', json={'online': True})
        r = self.client.post('/api/chat/send', json={'content': 'Tôi cần nhân viên tư vấn', 'mode': 'human'}).json()
        self.assertTrue(r['online'])
        self.assertEqual(r['messages'][-1]['role'], 'user')
        self.admin.post(f'/api/admin/chats/{r["id"]}/reply', json={'content': 'Dạ công ty đã nhận thông tin.'}).raise_for_status()
        self.assertEqual(self.client.get('/api/chat').json()['messages'][-1]['role'], 'admin')
        with connect() as c: c.execute('UPDATE sessions SET heartbeat=?', (time.time() - 60,))
        r = self.client.post('/api/chat/send', json={'content': 'Giá thế nào?', 'mode': 'human'}).json()
        self.assertFalse(r['online'])
        self.assertEqual(r['mode'], 'ai')
        self.assertEqual(r['messages'][-1]['role'], 'assistant')

    def test_real_ai_adapter_without_network(self):
        fake = MagicMock()
        fake.json.return_value = {'candidates': [{'content': {'parts': [{'text': 'Dạ, Nở hỗ trợ bạn chọn dịch vụ.'}]}}]}
        with patch.dict(os.environ, {'LLM_NAME': 'vertex', 'PROJECT_ID': 'test-project', 'LOCATION': 'us-central1', 'VERTEX_MODEL_NAME': 'test-model'}), patch('backend.logic.ai_ready', return_value=True), patch('backend.logic._vertex_access_token', return_value='test-token'), patch('backend.logic.httpx.post', return_value=fake) as call:
            r = self.client.post('/api/chat/send', json={'content': 'Xin chào', 'mode': 'ai'})
            self.assertEqual(r.json()['messages'][-1]['content'], 'Dạ, Nở hỗ trợ bạn chọn dịch vụ.')
            self.assertIn('systemInstruction', call.call_args.kwargs['json'])
            self.assertEqual(call.call_args.kwargs['headers']['Authorization'], 'Bearer test-token')
            self.assertIn('/v1/projects/test-project/locations/us-central1/', call.call_args.args[0])

    def test_order_lifecycle_requires_confirmed_price(self):
        order = self.create(mode='survey')
        path = f'/api/admin/orders/{order["id"]}'
        self.assertEqual(self.admin.patch(path, json={'status': 'paid'}).status_code, 422)
        self.assertEqual(self.admin.patch(path, json={'status': 'confirmed'}).status_code, 422)
        self.admin.patch(path + '/quote', json={'amount': 1500000, 'note': 'Giá đã thỏa thuận cho toàn bộ hạng mục.'}).raise_for_status()
        for status in ['surveying', 'confirmed', 'in_progress', 'completed', 'paid']:
            self.admin.patch(path, json={'status': status}).raise_for_status()
        self.assertEqual(self.client.get('/api/orders').json()[0]['status'], 'paid')
        self.assertEqual(self.admin.patch(path + '/quote', json={'amount': 5, 'note': 'Không được sửa'}).status_code, 422)

    def test_feedback_sends_admin_and_customer_email(self):
        feedback = {'name': 'Khách test', 'phone': '0901234567', 'email': 'customer@example.com', 'category': 'Thanh toán', 'rating': 3, 'subject': 'Kiểm tra thanh toán', 'content': 'Tôi cần kiểm tra giao dịch đã chuyển khoản.', 'transaction': 'TEST-123', 'consent': True}
        r = self.client.post('/api/feedback', json=feedback)
        self.assertEqual(r.status_code, 200)
        identifier = r.json()['id']
        admin_mail = self.admin.get('/api/admin/dashboard').json()['outbox'][0]
        self.assertEqual(admin_mail['kind'], 'feedback_admin')
        self.assertEqual(admin_mail['recipient'], 'admin@example.com')
        with patch.dict(os.environ, {'SMTP_HOST': 'smtp.example.com', 'SMTP_FROM': 'service@example.com'}), patch('backend.logic.smtplib.SMTP') as smtp:
            deliver_mail(admin_mail['id'])
            self.admin.patch(f'/api/admin/feedback/{identifier}', json={'status': 'resolved', 'reply': 'Đã kiểm tra và phản hồi cho khách.'}).raise_for_status()
            messages = [call.args[0] for call in smtp.return_value.__enter__.return_value.send_message.call_args_list]
            self.assertEqual(len(messages), 2)
            self.assertEqual(messages[0]['To'], 'admin@example.com')
            self.assertEqual(messages[0]['Reply-To'], 'customer@example.com')
            self.assertIn(identifier, messages[0]['Subject'])
            self.assertIn('TEST-123', messages[0].get_content())
            self.assertEqual(messages[1]['To'], 'customer@example.com')
            self.assertEqual(messages[1]['Reply-To'], 'admin@example.com')
            self.assertIn('Đã kiểm tra', messages[1].get_content())
            self.assertFalse(messages[0].is_multipart())
            self.assertFalse(messages[1].is_multipart())
        mine = self.client.get('/api/feedback').json()[0]
        self.assertEqual(mine['status'], 'resolved')
        self.assertIn('Đã kiểm tra', mine['reply'])
        mails = self.admin.get('/api/admin/dashboard').json()['outbox']
        self.assertEqual({mail['kind'] for mail in mails}, {'feedback_admin', 'feedback_customer'})
        self.assertEqual({mail['status'] for mail in mails}, {'sent'})

    def test_excel_formula_injection_is_plain_text(self):
        data = self.payload()
        data['name'] = '=HYPERLINK("https://example.com","click")'
        order = self.client.post('/api/orders', json=data).json()
        sheet = load_workbook(io.BytesIO(self.client.get(f'/api/orders/{order["id"]}/xlsx').content)).active
        self.assertEqual(sheet['B4'].data_type, 's')
        self.assertEqual(sheet['B4'].value, data['name'])

    def test_email_two_recipients_retry_and_attachment(self):
        order = self.create()
        mails = self.admin.get('/api/admin/dashboard').json()['outbox']
        self.assertEqual(len(mails), 2)
        self.assertEqual({m['status'] for m in mails}, {'pending'})
        with patch.dict(os.environ, {'SMTP_HOST': 'smtp.example.com', 'SMTP_FROM': 'service@example.com'}), patch('backend.logic.smtplib.SMTP') as smtp:
            for mail in mails: deliver_mail(mail['id'])
            messages = [call.args[0] for call in smtp.return_value.__enter__.return_value.send_message.call_args_list]
            self.assertEqual(len(messages), 2)
            self.assertEqual(len(list(messages[0].iter_attachments())), 1)
            deliver_mail(mails[0]['id'])
            self.assertEqual(smtp.return_value.__enter__.return_value.send_message.call_count, 2)
        self.assertEqual({m['status'] for m in self.client.get('/api/orders').json()[0]['mail']}, {'sent'})

    def test_admin_content_and_public_visibility(self):
        service = services()[0]
        service['active'] = False
        self.admin.put(f'/api/admin/services/{service["id"]}', json=service).raise_for_status()
        self.assertEqual(len(self.client.get('/api/bootstrap').json()['services']), 6)
        self.assertEqual(self.client.post('/api/orders', json=self.payload()).status_code, 404)
        body = {'title': 'Bài kiểm thử mới', 'category': 'Kiểm thử', 'excerpt': 'Nội dung mô tả', 'body': 'Nội dung kiểm thử dài hơn hai mươi ký tự.', 'image': '/images/hero.jpg', 'published': False}
        self.admin.put('/api/admin/blogs/new-blog', json=body).raise_for_status()
        self.assertFalse(any(b['id'] == 'new-blog' for b in self.client.get('/api/bootstrap').json()['blogs']))
        body['published'] = True
        self.admin.put('/api/admin/blogs/new-blog', json=body).raise_for_status()
        self.assertTrue(any(b['id'] == 'new-blog' for b in self.client.get('/api/bootstrap').json()['blogs']))

    def test_cross_origin_and_missing_csrf_header_rejected(self):
        self.assertEqual(self.client.post('/api/auth/logout', headers={'Origin': 'https://evil.example'}).status_code, 403)
        with TestClient(app) as stranger:
            self.assertEqual(stranger.post('/api/auth/logout').status_code, 403)

if __name__ == '__main__': unittest.main()
